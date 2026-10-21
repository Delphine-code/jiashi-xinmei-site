from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from pathlib import Path

out = Path('submission/2026_参赛提交材料工作版.docx')
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.75)
sec.right_margin = Inches(0.75)

PURPLE = '6B344D'
GOLD = 'B58A45'
INK = '30252B'
MUTED = '6E6962'
LINE = 'D8D0C5'
PALE = 'F7F3EC'
PALE_GOLD = 'F6EEDC'
PALE_PURPLE = 'F1E7EE'

styles = doc.styles
styles['Normal'].font.name = 'Microsoft YaHei'
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].font.color.rgb = RGBColor.from_string(INK)
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Normal'].paragraph_format.line_spacing = 1.25
for name, size, color in [('Title', 24, PURPLE), ('Heading 1', 16, PURPLE), ('Heading 2', 12.5, PURPLE), ('Heading 3', 11, GOLD)]:
    st = styles[name]
    st.font.name = 'Microsoft YaHei'
    st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(12 if name != 'Title' else 0)
    st.paragraph_format.space_after = Pt(6)

# Helpers
def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, color=LINE, size='6'):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag); borders.append(element)
        element.set(qn('w:val'), 'single'); element.set(qn('w:sz'), size); element.set(qn('w:space'), '0'); element.set(qn('w:color'), color)

def set_cell_text(cell, text, bold=False, color=INK, size=9.5):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text))
    r.bold = bold; r.font.name = 'Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei'); r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c, PURPLE); set_cell_text(c,h,True,'FFFFFF',9.2); set_cell_border(c,PURPLE,'8')
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            c=cells[i]; shade(c, 'FFFFFF' if ri%2 else PALE); set_cell_text(c,val,False,INK,9.1); set_cell_border(c)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def note(text, fill=PALE_GOLD):
    t=doc.add_table(rows=1, cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    c=t.cell(0,0); shade(c,fill); set_cell_border(c,GOLD,'8'); set_cell_text(c,text,False,INK,9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def bullet(text, level=0):
    p=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); p.paragraph_format.space_after=Pt(3); p.add_run(text)

def page_break(): doc.add_page_break()

# Cover
p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('2026 年参赛提交材料工作版')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('伽师新梅：戈壁绿洲上的紫色黄金'); r.bold=True; r.font.size=Pt(15); r.font.color.rgb=RGBColor.from_string(PURPLE)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('全国大学生数字媒体科技作品及创意竞赛'); r.font.size=Pt(10.5); r.font.color.rgb=RGBColor.from_string(MUTED)
doc.add_paragraph()
note('使用说明：本文件按官方提交规范整理为工作版。带“待核对 / 待补”的字段必须在提交前由团队确认，不应直接作为最终材料上传。')

# 1 system info
doc.add_heading('一、系统填报信息', level=1)
table(['字段','当前整理值','提交前动作'],[
['作品名称','伽师新梅：戈壁绿洲上的紫色黄金','与系统作品名称逐字核对'],
['参赛年份','2026 年','确认系统年份'],
['参赛学校','中国石油大学（北京）克拉玛依校区','核对学校全称'],
['参赛组别 / 赛道','本科；乡村振兴，科创赋能','以系统实际下拉选项为准'],
['指导教师','李兆美智（根据提交截图识别）','登录系统核对姓名与证件信息'],
['团队成员','郭婉红、张程、曾俊茹、李素琴（根据提交截图识别）','核对顺序、姓名和成员数量'],
['作品编号','406957（根据提交截图）','以系统当前编号为准'],
['AI 使用状态','使用 AI；需提交 AI 工具使用情况说明','与系统勾选项一致'],
['作品类型','多页交互式 Web 作品','与作品说明和展示视频保持一致'],
], [1.45,3.15,2.35])
note('截图中“待提交审核”显示需重点检查（7、8、9、10）项：作品展示视频、作品说明文档、素材来源说明、AI 工具使用说明等。')

# 2 work description
doc.add_heading('二、作品说明文案', level=1)
doc.add_heading('1. 项目定位', level=2)
doc.add_paragraph('《伽师新梅：戈壁绿洲上的紫色黄金》是一件以伽师新梅为主题的多页交互式数字叙事作品。作品从开屏地图进入，以“产地—风土—产业—科学—人物—品牌”的阅读路径，呈现一颗新梅如何在戈壁绿洲生长、被分选加工、进入市场，并转化为地方产业与品牌记忆。')
doc.add_heading('2. 主要内容', level=2)
bullet('产地定位：通过中国地图与新疆地图的层级交互，点击省域、新疆及地州市查看代表物产，突出伽师新梅的产地坐标。')
bullet('风土与生长：围绕日照、昼夜温差、昆仑雪水与土壤条件，结合四个月生长时间轴和糖度模拟器，解释风味形成过程。')
bullet('产业与流通：以种植面积、预计产量、冷链库容等关键数字呈现产业规模，并用自动冷链滑杆展示销售周期从约 20 天延展到 6 个月的逻辑。')
bullet('科学普及：以山梨糖醇、膳食纤维、花青素、钾、硒等成分为切入点，进行通俗化信息组织，并保留科学普及边界说明。')
bullet('人物与品牌：以种植、组织、经营、援疆、科研和流通等角色讲述产业协作，并展示梅梅 IP、包装设计、文旅立牌及 2026 中国新梅行业大会现场产品。')
doc.add_heading('3. 核心创新与实现', level=2)
table(['维度','说明'],[
['叙事结构','将地图定位、数据阅读、时间轴和人物故事组织为一条连续的作品阅读路径。'],
['交互方式','地图层级点击、悬停提示、梅梅图标预览、糖度模拟器、年份图表、自动冷链滑杆、人物卡片与图片放大浏览。'],
['技术实现','Astro 5 文件路由与页面组件；ECharts 5 地图与图表；原生 CSS 变量与响应式布局；静态资源本地化。'],
['视觉系统','以果紫、沙金、叶绿和米白纸张底构成统一色彩；采用标题衬线体与正文无衬线体，适配横版浏览器全屏观看、截图与录屏。'],
['真实成果','网页可在本地运行，主要页面与交互均已实现；提交前需用最终版本录制产品展示视频并检查所有链接。'],
],[1.4,5.55])

# 3 scoring alignment and defense notes
doc.add_heading('三、评分指标对照与答辩速记', level=1)
doc.add_paragraph('以下内容用于统一作品说明、产品展示视频和现场答辩口径。答辩时应以网页当前可操作的功能和已核验材料为准，不把概念设想表述为已完成成果。')
table(['评分重点','本作品对应证据','答辩时的一句话'],[
['选题价值与导向','以伽师新梅为入口，连接产地风土、产业流通、科学普及、人物协作与品牌表达，服务乡村振兴主题。','我们不是只展示一种水果，而是把产地知识、产业信息和地方品牌组织成可阅读、可验证的数字叙事。'],
['创新性与原创贡献','以中国—新疆—地州市三级地图作为入口，把地图、时间轴、数据图表、人物和品牌设计串成连续路径。','创新点在于把“找产地、看风土、看产业、看人物、看品牌”放进同一条交互阅读链。'],
['技术实现与规范','Astro 页面组件、ECharts 地图与图表、响应式 CSS、图片放大及拖动预览、自动冷链滑杆。','技术服务于信息表达：地图负责定位，图表负责变化，卡片负责判断，图片预览负责查看设计细节。'],
['完成度与可验证性','网页各主要路由可访问；地图、图表、滑杆、卡片与图片预览均可现场演示；来源页保留数据和素材说明。','评委可以通过网页和展示视频复核主要功能，提交材料与当前网站版本保持一致。'],
['应用与传播价值','适配横版浏览器全屏截图和录屏，可用于地方特产数字展示、文旅传播、产业科普和品牌提案。','作品将复杂的产地与产业信息转译成适合公众阅读的视觉体验。'],
], [1.45,3.5,2.0])
doc.add_heading('答辩高频问题', level=2)
table(['问题','建议回答'],[
['为什么选择伽师新梅？','伽师新梅同时具备鲜明产地、产业发展和品牌转化价值，适合用数字媒体把一颗果背后的地方知识讲清楚。'],
['作品的核心创新是什么？','不是单点做地图或数据，而是将产地定位、风土解释、产业变化、人物协作和品牌延展组织为一条完整的交互叙事。'],
['数据和图片如何保证可靠？','数据按来源、时间和口径整理；不同年份或口径不一致的内容不混用。大会现场图片单独标注来源，正式提交前完成授权与署名确认。'],
['AI 在作品中做了什么？','AI 主要辅助代码调试、文字润色、交互讨论和视觉方案比较；选题、信息架构、数据取舍、代码整合、事实核验和最终验收由团队完成。'],
['如果现场网络或页面异常怎么办？','准备最终静态网站压缩包、展示视频和关键页面截图；现场优先演示本地静态版本，并保留地图、数据和品牌页的录屏备份。'],
], [1.8,5.15])

# 4 video storyboard
page_break()
doc.add_heading('四、产品展示视频建议脚本', level=1)
note('官方规范重点要求展示真实产品效果、功能、交互体验和应用场景。建议使用 1920×1080、16:9、MP4；时长以系统当前要求为准。')
table(['时间','画面与操作','字幕 / 旁白建议'],[
['00:00—00:12','全屏进入 /entry，展示地图、主副标题与梅梅图标。','伽师新梅：戈壁绿洲上的紫色黄金。一颗果，从产地走向更远的地方。'],
['00:12—00:35','悬停省域显示名称；点击新疆进入新疆地图；点击地州市查看代表物产。','从中国找到新疆，再从新疆找到伽师；地图把产地关系变成可操作的阅读入口。'],
['00:35—01:05','进入 /terroir，展示四个风土条件，拖动生长时间轴。','光照、温差、雪水与土壤共同构成风味底色；四个月里，果实由青转紫。'],
['01:05—01:28','操作糖度模拟器，展示数值变化与实测区间。','把风土条件放进模拟器，观察糖度估算值变化，同时区分科普推演与实测口径。'],
['01:28—01:55','进入 /data，展示六个产业数字、年份图表和冷链自动滑杆。','产业规模、产量与品牌价值，回答这颗果如何长大；冷链把约二十天的采收窗口接成更长的销售周期。'],
['01:55—02:20','进入 /science，展示成分卡片、横向对比和免责声明。','从成分看药食同源，但不把科普内容替代为医疗建议。'],
['02:20—02:50','进入 /people，展开人物卡片，展示两张现场图片及来源注释。','一颗果的背后，是种植、技术、组织、流通等许多双手。'],
['02:50—03:20','进入 /brand，展示梅梅、礼盒、包装和大会产品展示；点击图片放大、缩放、拖动。','果实被带进包装、产品与文旅场景，品牌从一颗果变成可识别的视觉资产。'],
['03:20—03:35','回到 /about，展示来源、口径、素材说明；最后回到地图。','每个数字都有出处，每张图片都应完成来源和授权核对。'],
], [1.0,3.15,2.8])

# 5 AI declaration
doc.add_heading('五、AI 工具使用情况说明草案', level=1)
note('以下文字基于当前项目实际开发过程整理。提交前应由团队逐项核对工具名称、使用日期、具体模型、生成记录和最终贡献比例；不确定的内容不要直接提交。', PALE_PURPLE)
doc.add_heading('1. AI 使用状态', level=2)
doc.add_paragraph('本作品使用人工智能工具作为辅助工具，主要用于网页代码调试与迭代、交互方案讨论、文字润色和视觉排版建议。AI 未替代团队对作品主题、信息架构、数据取舍、页面结构、地图交互和最终成果的判断。')
doc.add_heading('2. AI 参与环节', level=2)
table(['环节','AI 辅助内容','团队完成的核心工作'],[
['网页开发','辅助分析 Astro 页面结构、CSS 样式、ECharts 地图 / 图表交互和浏览器问题。','确定页面架构、编写和修改代码、组织数据、调试交互、运行构建与页面验收。'],
['文字编辑','对页面标题、导语、图注和信息层级提出改写建议。','依据来源材料筛选事实、统一数据口径、确定最终表述并进行人工校对。'],
['视觉设计','辅助比较色彩、卡片、地图控件、图片排列和横版录屏布局。','确定果紫 / 沙金 / 叶绿视觉系统、最终排版、交互节奏和作品呈现方式。'],
['图片与素材','本材料中的大会现场实拍图片来自官方图集，不应表述为 AI 生成；梅梅及概念设计素材须按实际生成 / 创作记录填写。','确认图片来源、整理素材、裁切与页面适配，并负责参赛展示授权核对。'],
],[1.0,2.9,3.05])
doc.add_heading('3. 主要提示词 / 工作指令示例', level=2)
bullet('“检查网页地图部分的细节优化，统一地图与后续网页的视觉风格，保留主副标题。”')
bullet('“以专业网页设计和编辑视角，压缩说明性文字，强化风土—成熟—产业的阅读顺序。”')
bullet('“让冷链滑块自动播放，支持暂停 / 继续，适合横版浏览器录屏。”')
bullet('“检查图片比例、来源注释、放大预览、缩放与拖动交互。”')
doc.add_heading('4. 学术原创与人工修改说明', level=2)
doc.add_paragraph('作品的选题、叙事角度、页面信息架构、地图层级、数据呈现方式、交互流程、视觉系统和最终网页实现由参赛团队确定并完成。AI 输出仅作为辅助建议，团队对输出内容进行筛选、事实核验、改写、代码整合和功能测试；未经团队确认的 AI 输出不作为作品最终成果。')
doc.add_heading('5. 待补信息', level=2)
table(['待补项','填写内容'],[
['AI 工具名称与版本','待团队根据实际使用记录填写'],
['主要使用日期','待团队根据聊天 / 操作记录填写'],
['是否使用 AI 生成图片 / 音频 / 视频','请按实际情况勾选；若使用，逐项列明生成工具、用途、人工修改与来源'],
['AI 生成内容占比','按官方模板要求据实填写，不用“估计”替代实际记录'],
['生成记录 / 截图存档','建议保存关键提示词、输出版本、人工修改前后对比'],
], [2.4,4.55])

# 6 sources rights
doc.add_heading('六、素材来源与版权说明草案', level=1)
table(['素材','当前来源 / 用途','提交前必须完成'],[
['大会现场实拍图片','2026 中国新梅行业大会官方图集；用于人物页、品牌页产品展示等。','向主办方或摄影者确认参赛提交、公开展示、录屏和线上部署授权；确认署名方式并留存书面记录。'],
['地图边界数据','阿里云 DataV.GeoAtlas；用于中国、新疆及产地地图可视化示意。','保留来源说明，核对最新许可和竞赛展示要求。'],
['字体','Noto Sans SC / Noto Serif SC；用于网页标题与正文。','保留 SIL Open Font License 说明。'],
['梅梅 IP 及概念设计','页面中的 IP 图标、概念包装、挂件 / 玩偶等素材。','明确由团队原创、AI 辅助生成或外部素材；保存源文件、生成记录和人工修改记录。'],
['数据与文字','政府网站、公开报道、会议材料与团队整理材料。','逐条保留来源、日期和口径；预计值、调研材料与公开统计不得混写。'],
], [1.55,3.1,2.3])
doc.add_paragraph('建议最终在“关于与来源”页保留简洁来源说明，同时把完整授权证明、原图信息、截图和沟通记录放入参赛压缩包的“素材来源与版权说明”文件夹。', style='Normal')

# 7 checklist
doc.add_heading('七、提交前检查清单', level=1)
checks = [
'系统作品名称、学校、指导教师、团队成员、赛道和作品类型已逐项核对。',
'作品说明文档与网页当前版本一致，未把计划功能写成已实现功能。',
'产品展示视频已按 16:9、1920×1080、MP4 导出，网页交互和真实页面均可看见。',
'作品截图不少于官方要求数量，优先选择地图、风土、数据、科学、人物、品牌和来源页的代表画面。',
'作品截图和视频中不出现本地路径、调试面板、无关浏览器标签、个人账号或未授权素材。',
'大会现场实拍图片已获得参赛与公开展示授权，并按要求署名。',
'AI 使用说明、提示词记录、人工修改记录和生成素材来源已整理归档。',
'所有来源说明、授权文件、作品说明、展示视频、截图和其他材料按官方压缩包结构命名。',
'本地运行和公网部署均已测试；主要路由刷新不 404，地图、图片、视频和链接可正常打开。',
'最终压缩包解压后重新检查，视频可播放，PDF 可打开，文件名和提交系统要求一致。',
]
for c in checks: bullet('□ ' + c)
note('当前最优先的三项：①确认官方图集授权与署名；②补齐 AI 工具名称、日期和生成记录；③录制并复核最终产品展示视频。')

# footer
for section in doc.sections:
    footer=section.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=footer.add_run('2026 参赛提交材料工作版  ·  伽师新梅：戈壁绿洲上的紫色黄金')
    run.font.size=Pt(8); run.font.color.rgb=RGBColor.from_string(MUTED)

# core properties
doc.core_properties.title = '2026参赛提交材料工作版'
doc.core_properties.subject = '伽师新梅数字媒体作品参赛材料'
doc.core_properties.author = '参赛团队'
doc.core_properties.keywords = '伽师新梅, 数字媒体, 参赛材料, AI使用说明'
doc.save(out)
print(out.resolve())
